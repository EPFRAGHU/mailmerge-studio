"""
MailMerge Studio — a friendly desktop mail-merge app.
Import an Excel/CSV contact list, compose a rich email with {{merge_tags}},
attach files (global or per-row), and send from your own Gmail/Outlook/SMTP.
"""
import json
import os
import re
import smtplib
import ssl
import threading
import time
import uuid
from datetime import datetime
from email.message import EmailMessage
from email.utils import formataddr, parseaddr

from flask import Flask, jsonify, render_template, request, send_from_directory
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("DATA_DIR") or os.path.join(BASE, "data")
UPLOADS = os.path.join(DATA, "uploads")
ATTACH = os.path.join(DATA, "attachments")
os.makedirs(UPLOADS, exist_ok=True)
os.makedirs(ATTACH, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024

STATE = {
    "rows": [],            # contact rows (dicts)
    "columns": [],         # column names
    "source_file": None,
    "config": {},          # smtp settings + campaign settings
    "campaign": None,      # active send job
}

CONFIG_PATH = os.path.join(DATA, "config.json")
LOG_PATH = os.path.join(DATA, "send_log.json")


def load_state():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r", encoding="utf-8") as f:
                STATE["config"] = json.load(f)
        except Exception:
            pass


def save_state():
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(STATE["config"], f, indent=2, ensure_ascii=False)


def save_log(entries):
    with open(LOG_PATH, "w", encoding="utf-8") as f:
        json.dump(entries, f, indent=2, ensure_ascii=False)


DAILY_PATH = os.path.join(DATA, "daily_stats.json")


def bump_daily_stats(entries):
    """Accumulate sent/failed counts per calendar day (persistent across restarts)."""
    stats = {}
    if os.path.exists(DAILY_PATH):
        try:
            with open(DAILY_PATH, "r", encoding="utf-8") as f:
                stats = json.load(f)
        except Exception:
            stats = {}
    for e in entries:
        day = str(e.get("at", ""))[:10]
        if not day:
            continue
        d = stats.setdefault(day, {"sent": 0, "failed": 0})
        d["sent" if e.get("status") == "sent" else "failed"] += 1
    with open(DAILY_PATH, "w", encoding="utf-8") as f:
        json.dump(stats, f, indent=2)


def read_daily_stats():
    if os.path.exists(DAILY_PATH):
        try:
            with open(DAILY_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def read_log():
    if os.path.exists(LOG_PATH):
        try:
            with open(LOG_PATH, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []


load_state()

# Auto-restore the last used mailing list after a restart
_last = STATE["config"].get("last_list")
if _last:
    _lpath = os.path.join(DATA, "lists", re.sub(r"[^A-Za-z0-9_-]+", "_", str(_last).strip())[:60] + ".json")
    if os.path.exists(_lpath):
        try:
            with open(_lpath, "r", encoding="utf-8") as _f:
                _d = json.load(_f)
            STATE["rows"], STATE["columns"] = _d.get("rows", []), _d.get("columns", [])
            STATE["source_file"] = _d.get("source", _last)
        except Exception:
            pass

# ---------------- persistent lists & templates ----------------
LISTS_DIR = os.path.join(DATA, "lists")
TPL_DIR = os.path.join(DATA, "templates")
os.makedirs(LISTS_DIR, exist_ok=True)
os.makedirs(TPL_DIR, exist_ok=True)


def slugify(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", s.strip())[:60] or "list"


class SafeJSONEncoder(json.JSONEncoder):
    def default(self, o):
        try:
            import datetime as _dt
            if isinstance(o, (_dt.datetime, _dt.date)):
                return o.strftime("%d/%m/%Y")
        except Exception:
            pass
        return str(o)


def list_meta(path, data):
    return {"name": data.get("name", os.path.basename(path)), "created": data.get("created", ""),
            "count": len(data.get("rows", [])), "columns": len(data.get("columns", []))}


@app.route("/api/lists")
def api_lists():
    out = []
    for f in sorted(os.listdir(LISTS_DIR)):
        try:
            with open(os.path.join(LISTS_DIR, f), "r", encoding="utf-8") as fh:
                out.append(list_meta(f, json.load(fh)))
        except Exception:
            pass
    return jsonify(lists=out)


@app.route("/api/lists/save", methods=["POST"])
def api_lists_save():
    name = (request.json.get("name") or "").strip() or "My List"
    data = {"name": name, "created": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "columns": STATE["columns"], "rows": STATE["rows"], "source": STATE["source_file"]}
    with open(os.path.join(LISTS_DIR, slugify(name) + ".json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, cls=SafeJSONEncoder)
    STATE["config"]["last_list"] = name
    save_state()
    return jsonify(ok=True)


@app.route("/api/lists/load", methods=["POST"])
def api_lists_load():
    name = request.json.get("name", "")
    path = os.path.join(LISTS_DIR, slugify(name) + ".json")
    if not os.path.exists(path):
        return jsonify(error="List not found"), 404
    with open(path, "r", encoding="utf-8") as f:
        d = json.load(f)
    STATE["rows"], STATE["columns"] = d["rows"], d["columns"]
    STATE["source_file"] = d.get("source", name)
    EMAIL_COL["value"] = find_email_column(STATE["columns"])
    STATE["config"]["last_list"] = name
    save_state()
    return jsonify(columns=STATE["columns"], rows=STATE["rows"], total=len(STATE["rows"]),
                   source=STATE["source_file"], emailColumn=EMAIL_COL["value"])


@app.route("/api/lists/delete", methods=["POST"])
def api_lists_delete():
    name = request.json.get("name", "")
    path = os.path.join(LISTS_DIR, slugify(name) + ".json")
    if os.path.exists(path):
        os.remove(path)
    return jsonify(ok=True)


@app.route("/api/templates")
def api_templates():
    out = []
    for f in sorted(os.listdir(TPL_DIR)):
        try:
            with open(os.path.join(TPL_DIR, f), "r", encoding="utf-8") as fh:
                d = json.load(fh)
            out.append({"name": d.get("name", f), "list": d.get("list", ""), "updated": d.get("updated", "")})
        except Exception:
            pass
    return jsonify(templates=out)


@app.route("/api/templates/save", methods=["POST"])
def api_templates_save():
    d = request.json or {}
    name = (d.get("name") or "").strip() or "Untitled template"
    data = {"name": name, "subject": d.get("subject", ""), "html": d.get("html", ""),
            "list": d.get("list", ""), "updated": datetime.now().strftime("%Y-%m-%d %H:%M")}
    with open(os.path.join(TPL_DIR, slugify(name) + ".json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, cls=SafeJSONEncoder)
    return jsonify(ok=True)


@app.route("/api/templates/load", methods=["POST"])
def api_templates_load():
    name = request.json.get("name", "")
    path = os.path.join(TPL_DIR, slugify(name) + ".json")
    if not os.path.exists(path):
        return jsonify(error="Template not found"), 404
    with open(path, "r", encoding="utf-8") as f:
        return jsonify(json.load(f))


@app.route("/api/templates/delete", methods=["POST"])
def api_templates_delete():
    name = request.json.get("name", "")
    path = os.path.join(TPL_DIR, slugify(name) + ".json")
    if os.path.exists(path):
        os.remove(path)
    return jsonify(ok=True)


@app.route("/api/state")
def api_state():
    """Current in-memory workspace so a page refresh doesn't lose anything."""
    return jsonify(columns=STATE["columns"], rows=STATE["rows"], total=len(STATE["rows"]),
                   source=STATE["source_file"], emailColumn=EMAIL_COL["value"],
                   lastList=STATE["config"].get("last_list", ""))

# ---------------- data import ----------------


def parse_csv(path):
    import csv
    rows, cols = [], []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except Exception:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        all_rows = list(reader)
    if not all_rows:
        return [], []
    cols = [c.strip() for c in all_rows[0]]
    for r in all_rows[1:]:
        if not any(str(x).strip() for x in r):
            continue
        rows.append({cols[i] if i < len(cols) else f"Column{i+1}": v for i, v in enumerate(r)})
    return rows, cols


def parse_xlsx(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    all_rows = [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
    while all_rows and not any(str(x).strip() for x in all_rows[0]):
        all_rows.pop(0)
    if not all_rows:
        return [], []
    cols = [str(c).strip() if str(c).strip() else f"Column{i+1}" for i, c in enumerate(all_rows[0])]
    rows = []
    for r in all_rows[1:]:
        if not any(str(x).strip() for x in r):
            continue
        row = {}
        for i, c in enumerate(cols):
            row[c] = "" if i >= len(r) else r[i]
        rows.append(row)
    return rows, cols


@app.route("/api/import", methods=["POST"])
def api_import():
    file = request.files.get("file")
    if not file:
        return jsonify(error="No file provided"), 400
    name = file.filename or "data"
    ext = os.path.splitext(name)[1].lower()
    path = os.path.join(UPLOADS, f"{uuid.uuid4().hex}{ext}")
    file.save(path)
    try:
        if ext in (".csv", ".tsv", ".txt"):
            rows, cols = parse_csv(path)
        elif ext in (".xlsx", ".xlsm"):
            rows, cols = parse_xlsx(path)
        else:
            return jsonify(error="Unsupported file type. Use .xlsx, .csv or .tsv"), 400
    except Exception as e:
        return jsonify(error=f"Could not read file: {e}"), 400
    if not rows or not cols:
        return jsonify(error="The file appears to be empty"), 400
    STATE["rows"], STATE["columns"] = rows, cols
    STATE["source_file"] = name
    EMAIL_COL["value"] = find_email_column(cols)
    return jsonify(columns=cols, rows=rows[:50], total=len(rows), source=name,
                   emailColumn=EMAIL_COL["value"])


# ---------------- merge helpers ----------------

TAG_RE = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")


def render(text, row):
    def sub(m):
        key = m.group(1)
        val = row.get(key, "")
        return "" if val is None else str(val)
    return TAG_RE.sub(sub, text or "")


def find_email_column(cols):
    """Pick the most likely email column from a list of header names."""
    for exact in ("email", "email id", "email_id", "emailid", "e-mail", "mail id", "mail"):
        for c in cols:
            if c.strip().lower() == exact:
                return c
    for c in cols:
        if "email" in c.strip().lower():
            return c
    return None


EMAIL_COL = {"value": None}


def valid_email(a):
    try:
        addr = parseaddr(str(a))[1]
        return "@" in addr and "." in addr.split("@")[-1]
    except Exception:
        return False


# ---------------- attachments ----------------


@app.route("/api/attachments", methods=["POST"])
def api_upload_attachment():
    files = request.files.getlist("files")
    saved = []
    for f in files:
        name = f.filename or f"file_{uuid.uuid4().hex}"
        dest = os.path.join(ATTACH, name)
        f.save(dest)
        saved.append(name)
    return jsonify(saved=saved)


@app.route("/api/attachments", methods=["GET"])
def api_list_attachments():
    return jsonify(files=sorted(os.listdir(ATTACH)))


@app.route("/api/attachments/delete", methods=["POST"])
def api_delete_attachment():
    name = request.json.get("name", "")
    safe = os.path.basename(name)
    p = os.path.join(ATTACH, safe)
    if os.path.exists(p):
        os.remove(p)
    return jsonify(ok=True)


# ---------------- sending engine ----------------

class Campaign(threading.Thread):
    def __init__(self, payload):
        super().__init__(daemon=True)
        self.payload = payload
        self.id = uuid.uuid4().hex[:8]
        self.stop_flag = threading.Event()
        self.log = []
        self.status = "running"
        self.sent = self.failed = self.total = 0

    def run(self):
        try:
            self._run()
        except Exception as e:
            self.status = "error"
            self.log.append({"to": "-", "status": "failed", "error": f"Campaign crashed: {e}", "at": datetime.now().isoformat(timespec="seconds")})
        finally:
            save_log(self.log)
            if self.status == "running":
                self.status = "done"

    def _run(self):
        p = self.payload
        subject_t = p.get("subject", "")
        html = p.get("html", "")
        global_attach = p.get("attachments", [])
        attach_col = p.get("attachmentColumn")
        throttle = max(0.5, float(p.get("delaySeconds", 3)))
        batch = max(1, int(p.get("batchSize", 40)))
        pause = max(30, int(p.get("batchPauseSeconds", 120)))
        dry = bool(p.get("testMode", False))
        test_to = p.get("testRecipient", "")

        cfg = STATE["config"].get("smtp", {})
        host = cfg.get("host", "")
        port = int(cfg.get("port", 587))
        user = cfg.get("user", "")
        pwd = cfg.get("password", "")
        use_tls = cfg.get("security", "starttls") == "starttls"
        use_ssl = cfg.get("security", "starttls") == "ssl"
        from_name = cfg.get("fromName", "")
        from_addr = cfg.get("fromEmail") or user
        reply_to = cfg.get("replyTo", "")

        rows = STATE["rows"] if not dry else [{"Email": test_to or "test@example.com"}]
        # filter to selected rows if provided
        selected = p.get("selectedRows")
        if selected is not None and not dry:
            rows = [STATE["rows"][i] for i in selected if 0 <= i < len(STATE["rows"])]

        self.total = len(rows)
        count = 0
        for row in rows:
            if self.stop_flag.is_set():
                self.status = "stopped"
                break
            ecol = EMAIL_COL["value"]
            to = str(row.get(ecol) or row.get("Email") or row.get("email") or "").strip()
            if dry:
                to = test_to or "test@example.com"
            if not valid_email(to):
                self.log.append({"to": to, "status": "failed", "error": "Invalid/missing email", "at": datetime.now().isoformat(timespec="seconds")})
                self.failed += 1
                count += 1
                continue
            subject = render(subject_t, row)
            body = render(html, row)
            files = list(global_attach)
            if attach_col:
                extra = str(row.get(attach_col, "") or "").strip()
                if extra:
                    files += [x.strip() for x in extra.split(";") if x.strip()]
            ok, err = True, ""
            if not dry:
                try:
                    msg = EmailMessage()
                    msg["From"] = formataddr((from_name or from_addr, from_addr))
                    msg["To"] = to
                    if reply_to:
                        msg["Reply-To"] = reply_to
                    msg["Subject"] = subject
                    msg.set_content("Please view this email in an HTML-capable client.")
                    msg.add_alternative(body, subtype="html")
                    for fn in files:
                        fp = os.path.join(ATTACH, os.path.basename(fn))
                        if os.path.exists(fp):
                            with open(fp, "rb") as fh:
                                msg.add_attachment(fh.read(), maintype="application", subtype="octet-stream", filename=os.path.basename(fn))
                    if use_ssl:
                        with smtplib.SMTP_SSL(host, port, timeout=30, context=ssl.create_default_context()) as s:
                            s.login(user, pwd)
                            s.send_message(msg)
                    else:
                        with smtplib.SMTP(host, port, timeout=30) as s:
                            s.ehlo()
                            if use_tls:
                                s.starttls(context=ssl.create_default_context())
                            s.login(user, pwd)
                            s.send_message(msg)
                except Exception as e:
                    ok, err = False, str(e)
            entry = {"to": to, "status": "sent" if ok else "failed", "error": err, "at": datetime.now().isoformat(timespec="seconds")}
            self.log.append(entry)
            if ok:
                self.sent += 1
            else:
                self.failed += 1
            count += 1
            time.sleep(throttle)
            if not dry and count % batch == 0 and count < self.total:
                for _ in range(pause):
                    if self.stop_flag.is_set():
                        break
                    time.sleep(1)
        if self.status != "stopped":
            self.status = "done"
        save_log(self.log)
        if not dry:
            bump_daily_stats(self.log)


@app.route("/api/send", methods=["POST"])
def api_send():
    if STATE["campaign"] and STATE["campaign"].is_alive():
        return jsonify(error="A campaign is already running"), 409
    camp = Campaign(request.json)
    STATE["campaign"] = camp
    camp.start()
    return jsonify(id=camp.id)


@app.route("/api/stop", methods=["POST"])
def api_stop():
    c = STATE["campaign"]
    if c:
        c.stop_flag.set()
    return jsonify(ok=True)


@app.route("/api/progress")
def api_progress():
    c = STATE["campaign"]
    if not c:
        prev = read_log()
        return jsonify(running=False, log=prev[-200:], sent=sum(1 for x in prev if x["status"] == "sent"),
                       failed=sum(1 for x in prev if x["status"] == "failed"), total=len(prev), status="idle")
    return jsonify(running=c.is_alive(), status=c.status, sent=c.sent, failed=c.failed,
                   total=c.total, log=c.log[-200:])


@app.route("/api/log/export")
def api_log_export():
    return send_from_directory(DATA, "send_log.json", as_attachment=True, download_name="send_report.json")


DAILY_LIMIT = 500  # Gmail consumer cap; adjust if your mailbox differs


@app.route("/api/daily")
def api_daily():
    stats = read_daily_stats()
    today = datetime.now().strftime("%Y-%m-%d")
    t = stats.get(today, {"sent": 0, "failed": 0})
    days = sorted(stats.items(), reverse=True)[:30]
    return jsonify(limit=DAILY_LIMIT, today=t, days=days)


# ---------------- config ----------------

@app.route("/api/config", methods=["GET", "POST"])
def api_config():
    if request.method == "POST":
        c = request.json or {}
        # Preserve the stored password when the update doesn't include one
        old_pwd = STATE["config"].get("smtp", {}).get("password", "")
        new_smtp = c.get("smtp")
        if isinstance(new_smtp, dict) and not new_smtp.get("password"):
            new_smtp["password"] = old_pwd
        STATE["config"].update(c)
        save_state()
        return jsonify(ok=True)
    cfg = json.loads(json.dumps(STATE["config"]))
    if "smtp" in cfg:
        cfg["smtp"]["password"] = "********" if cfg["smtp"].get("password") else ""
    return jsonify(cfg)


@app.route("/api/config/save-password", methods=["POST"])
def api_save_password():
    pwd = request.json.get("password", "")
    STATE["config"].setdefault("smtp", {})["password"] = pwd
    save_state()
    return jsonify(ok=True)


@app.route("/api/test-connection", methods=["POST"])
def api_test_connection():
    """Verify SMTP login only — sends nothing."""
    cfg = STATE["config"].get("smtp", {})
    host = cfg.get("host", "")
    port = int(cfg.get("port", 587))
    user = cfg.get("user", "")
    pwd = cfg.get("password", "")
    sec = cfg.get("security", "starttls")
    if not host or not user or not pwd:
        return jsonify(ok=False, error="Save your connection (host, email and password) first")
    try:
        if sec == "ssl":
            with smtplib.SMTP_SSL(host, port, timeout=15, context=ssl.create_default_context()) as s:
                s.login(user, pwd)
        else:
            with smtplib.SMTP(host, port, timeout=15) as s:
                s.ehlo()
                if sec == "starttls":
                    s.starttls(context=ssl.create_default_context())
                s.login(user, pwd)
        return jsonify(ok=True)
    except smtplib.SMTPAuthenticationError:
        return jsonify(ok=False, error="Authentication rejected — check the app password (535 error)")
    except Exception as e:
        return jsonify(ok=False, error=f"Could not reach server: {e}")


PRESETS = {
    "gmail": {"host": "smtp.gmail.com", "port": 587, "security": "starttls"},
    "outlook": {"host": "smtp.office365.com", "port": 587, "security": "starttls"},
    "yahoo": {"host": "smtp.mail.yahoo.com", "port": 587, "security": "starttls"},
    "zoho": {"host": "smtp.zoho.com", "port": 587, "security": "starttls"},
    "icloud": {"host": "smtp.mail.me.com", "port": 587, "security": "starttls"},
    "custom": {"host": "", "port": 587, "security": "starttls"},
}


@app.route("/api/presets")
def api_presets():
    return jsonify(PRESETS)


@app.route("/")
def index():
    # Serve as a raw file: the page contains {{merge_tag}} syntax that
    # Jinja2 templating would misinterpret.
    return send_from_directory(os.path.join(BASE, "templates"), "index.html")


if __name__ == "__main__":
    print("\n  MailMerge Studio is running →  http://127.0.0.1:5000\n")
    app.run(host="127.0.0.1", port=5000, debug=False)